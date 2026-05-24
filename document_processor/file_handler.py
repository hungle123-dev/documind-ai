import hashlib
import os
import pickle
from datetime import datetime, timedelta
from pathlib import Path
from typing import Any

from langchain_core.documents import Document

from config.settings import settings
from utils.language import detect_language
from utils.logging import logger


class DocumentProcessor:
    CACHE_SCHEMA_VERSION = 7

    def __init__(self, cache_dir: Path | str | None = None) -> None:
        self.cache_dir = Path(cache_dir or settings.CACHE_DIR)
        self.cache_dir.mkdir(parents=True, exist_ok=True)

    def validate_files(self, files: list[Any]) -> None:
        total_size = 0
        for file in files:
            path = self._file_path(file)
            if path.suffix.lower() not in settings.ALLOWED_TYPES:
                raise ValueError(f"Unsupported file type: {path.suffix}")
            size = path.stat().st_size
            if size > settings.MAX_FILE_SIZE:
                limit_mb = settings.MAX_FILE_SIZE // 1024 // 1024
                raise ValueError(f"File {path.name} exceeds {limit_mb}MB limit")
            total_size += size

        if total_size > settings.MAX_TOTAL_SIZE:
            limit_mb = settings.MAX_TOTAL_SIZE // 1024 // 1024
            raise ValueError(f"Total size exceeds {limit_mb}MB limit")

    def process(self, files: list[Any]) -> list[Document]:
        self.validate_files(files)
        chunks: list[Document] = []
        seen_content_hashes: set[str] = set()

        for file in files:
            path = self._file_path(file)
            file_hash = self._generate_hash(path.read_bytes())
            cache_path = self.cache_dir / f"{file_hash}.pkl"

            if self._is_cache_valid(cache_path):
                logger.info(f"Loading cached chunks for {path.name}")
                file_chunks = self._load_from_cache(cache_path)
            else:
                logger.info(f"Processing document {path.name}")
                file_chunks = self._process_file(path, file_hash)
                self._save_to_cache(file_chunks, cache_path)

            for chunk in file_chunks:
                chunk_hash = self._generate_hash(chunk.page_content.encode("utf-8"))
                if chunk_hash not in seen_content_hashes:
                    chunks.append(chunk)
                    seen_content_hashes.add(chunk_hash)

        logger.info(f"Total unique chunks: {len(chunks)}")
        return chunks

    def _process_file(self, path: Path, file_hash: str) -> list[Document]:
        suffix = path.suffix.lower()
        if suffix == ".pdf":
            return self._process_pdf(path, file_hash)
        elif suffix == ".docx":
            text = self._parse_docx(path)
        elif suffix in {".txt", ".md"}:
            text = self._parse_text(path)
        else:
            raise ValueError(f"Unsupported file type: {suffix}")

        if not text.strip():
            logger.warning(f"No readable text extracted from {path.name}")
            return []

        return self._build_chunks(text, source=path.name, file_hash=file_hash)

    def _process_pdf(self, path: Path, file_hash: str) -> list[Document]:
        pages = self._parse_pdf(path)
        full_text = "\n\n".join(text for _, text in pages)
        if not full_text.strip():
            logger.warning(f"No readable text extracted from {path.name}")
            return []

        language = detect_language(full_text)
        documents: list[Document] = []
        for page_number, text in pages:
            documents.extend(
                self._build_chunks(
                    text,
                    source=path.name,
                    file_hash=file_hash,
                    page=page_number,
                    language=language,
                    start_index=len(documents),
                )
            )
        return documents

    def _parse_pdf(self, path: Path) -> list[tuple[int, str]]:
        try:
            import pymupdf4llm
        except ImportError as exc:
            raise RuntimeError("pymupdf4llm is required to parse PDF files.") from exc

        page_chunks = pymupdf4llm.to_markdown(str(path), page_chunks=True)
        pages = [
            (int(chunk["metadata"]["page_number"]), str(chunk["text"]))
            for chunk in page_chunks
            if str(chunk.get("text", "")).strip()
        ]
        if len("\n".join(text for _, text in pages).strip()) < 100:
            logger.warning(f"{path.name} appears to be scanned or has no extractable text.")
            return []
        return pages

    def _parse_docx(self, path: Path) -> str:
        from docx import Document as DocxDocument

        document = DocxDocument(str(path))
        parts: list[str] = []
        for paragraph in document.paragraphs:
            text = paragraph.text.strip()
            if text:
                parts.append(text)
        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                if any(cells):
                    parts.append(" | ".join(cells))
        return "\n\n".join(parts)

    def _parse_text(self, path: Path) -> str:
        return path.read_text(encoding="utf-8")

    def _build_chunks(
        self,
        text: str,
        source: str,
        file_hash: str,
        page: int | None = None,
        language: str | None = None,
        start_index: int = 0,
    ) -> list[Document]:
        document_language = language or detect_language(text)
        raw_chunks = self._split_text(text)
        documents: list[Document] = []

        for index, chunk in enumerate(raw_chunks):
            documents.append(
                Document(
                    page_content=chunk,
                    metadata={
                        "source": source,
                        "file_hash": file_hash,
                        "processor_schema_version": self.CACHE_SCHEMA_VERSION,
                        "chunk_index": start_index + index,
                        "page": page,
                        "section": self._nearest_section(chunk),
                        "language": document_language,
                    },
                )
            )
        return documents

    def _split_text(self, text: str, chunk_size: int = 800, chunk_overlap: int = 200) -> list[str]:
        cleaned = text.strip()
        if not cleaned:
            return []
        chunks: list[str] = []
        for block, is_table in self._markdown_blocks(cleaned):
            if is_table:
                if chunks and self._is_table_caption(chunks[-1]):
                    block = chunks.pop() + "\n" + block
                chunks.append(self._linearize_markdown_table(block))
            else:
                chunks.extend(self._split_plain_text(block, chunk_size, chunk_overlap))
        return [chunk for chunk in chunks if chunk.strip()]

    def _is_table_caption(self, text: str) -> bool:
        stripped = text.strip().lower()
        return stripped.startswith("table ") or stripped.startswith("table:")

    def _linearize_markdown_table(self, table: str) -> str:
        rows = [
            self._parse_markdown_table_row(line)
            for line in table.splitlines()
            if line.strip().startswith("|") and line.strip().endswith("|")
        ]
        rows = [row for row in rows if row and not self._is_markdown_separator(row)]
        if len(rows) < 2:
            return table.strip()

        header = rows[0]
        linearized: list[str] = []
        for row in rows[1:]:
            linearized.extend(self._linearize_table_row(header, row))
        if not linearized:
            return table.strip()
        return table.strip() + "\n\nLinearized table rows:\n" + "\n".join(linearized)

    def _parse_markdown_table_row(self, line: str) -> list[str]:
        stripped = line.strip().strip("|")
        return [cell.strip() for cell in stripped.split("|")]

    def _is_markdown_separator(self, row: list[str]) -> bool:
        return all(cell and set(cell) <= {"-", ":", " "} for cell in row)

    def _cell_parts(self, cell: str) -> list[str]:
        return [part.strip().strip("*_ ") for part in cell.split("<br>") if part.strip().strip("*_ ")]

    def _linearize_table_row(self, header: list[str], row: list[str]) -> list[str]:
        entity_parts = self._cell_parts(row[0]) if row else []
        if not entity_parts:
            return []
        lines: list[str] = []
        for entity_index, entity in enumerate(entity_parts):
            attributes: list[str] = []
            for column_index, cell in enumerate(row[1:], start=1):
                header_text = header[column_index] if column_index < len(header) else f"Column {column_index + 1}"
                attributes.extend(
                    self._linearize_table_cell(
                        header_text,
                        cell,
                        entity_index,
                        len(entity_parts),
                    )
                )
            if attributes:
                lines.append(f"{entity}: {'; '.join(attributes)}")
        return lines

    def _linearize_table_cell(
        self,
        header_text: str,
        cell: str,
        entity_index: int,
        entity_count: int,
    ) -> list[str]:
        header_parts = self._cell_parts(header_text)
        values = self._cell_parts(cell)
        if not values:
            return []

        if entity_count > 1 and len(values) % entity_count == 0:
            values_per_entity = len(values) // entity_count
            start = entity_index * values_per_entity
            entity_values = values[start : start + values_per_entity]
        elif entity_index < len(values):
            entity_values = [values[entity_index]]
        else:
            return []

        label = header_parts[0] if header_parts else header_text.strip()
        sublabels = header_parts[1:]
        if len(sublabels) >= len(entity_values):
            return [
                f"{label} {sublabels[index]} = {value}"
                for index, value in enumerate(entity_values)
            ]
        if len(entity_values) == 1:
            return [f"{label} = {entity_values[0]}"]
        return [f"{label} value {index + 1} = {value}" for index, value in enumerate(entity_values)]

    def _markdown_blocks(self, text: str) -> list[tuple[str, bool]]:
        blocks: list[tuple[str, bool]] = []
        text_lines: list[str] = []
        table_lines: list[str] = []

        def flush_text() -> None:
            if text_lines:
                block = "\n".join(text_lines).strip()
                if block:
                    blocks.append((block, False))
                text_lines.clear()

        def flush_table() -> None:
            if table_lines:
                blocks.append(("\n".join(table_lines).strip(), True))
                table_lines.clear()

        for line in text.splitlines():
            stripped = line.strip()
            if stripped.startswith("|") and stripped.endswith("|"):
                flush_text()
                table_lines.append(line)
            else:
                flush_table()
                text_lines.append(line)

        flush_table()
        flush_text()
        return blocks

    def _split_plain_text(
        self,
        text: str,
        chunk_size: int = 800,
        chunk_overlap: int = 200,
    ) -> list[str]:
        cleaned = text.strip()
        if not cleaned:
            return []
        chunks: list[str] = []
        start = 0
        while start < len(cleaned):
            end = min(start + chunk_size, len(cleaned))
            if end < len(cleaned):
                split_at = max(
                    cleaned.rfind("\n\n", start, end),
                    cleaned.rfind("\n", start, end),
                    cleaned.rfind(". ", start, end),
                    cleaned.rfind(" ", start, end),
                )
                if split_at > start + chunk_size // 2:
                    end = split_at + 1
            chunk = cleaned[start:end].strip()
            if chunk:
                chunks.append(chunk)
            if end >= len(cleaned):
                break
            next_start = self._sentence_overlap_start(cleaned, end, chunk_overlap)
            if next_start <= start:
                next_start = max(start + 1, end - chunk_overlap)
            start = next_start
        return chunks

    def _sentence_overlap_start(self, text: str, end: int, chunk_overlap: int) -> int:
        overlap_start = max(0, end - chunk_overlap)
        search_start = max(0, overlap_start - chunk_overlap)
        sentence_starts = [
            index + 2
            for index in range(search_start, end)
            if text[index : index + 2] in {". ", "? ", "! "}
        ]
        candidates = [index for index in sentence_starts if index <= overlap_start]
        if candidates:
            return candidates[-1]
        return overlap_start

    def _nearest_section(self, chunk: str) -> str:
        for line in chunk.splitlines():
            stripped = line.strip()
            if stripped.startswith("#"):
                return stripped.lstrip("#").strip() or "N/A"
        first_line = chunk.splitlines()[0].strip() if chunk.splitlines() else ""
        return first_line[:80] if first_line else "N/A"

    def _file_path(self, file: Any) -> Path:
        if isinstance(file, (str, os.PathLike)):
            return Path(file)
        return Path(file.name)

    def _generate_hash(self, content: bytes) -> str:
        return hashlib.sha256(content).hexdigest()

    def _save_to_cache(self, chunks: list[Document], cache_path: Path) -> None:
        cache_path.parent.mkdir(parents=True, exist_ok=True)
        with cache_path.open("wb") as file:
            pickle.dump(
                {
                    "schema_version": self.CACHE_SCHEMA_VERSION,
                    "timestamp": datetime.now().timestamp(),
                    "chunks": chunks,
                },
                file,
            )

    def _load_from_cache(self, cache_path: Path) -> list[Document]:
        with cache_path.open("rb") as file:
            data = pickle.load(file)
        return data["chunks"]

    def _is_cache_valid(self, cache_path: Path) -> bool:
        if not cache_path.exists():
            return False
        cache_age = datetime.now() - datetime.fromtimestamp(cache_path.stat().st_mtime)
        if cache_age >= timedelta(days=settings.CACHE_EXPIRE_DAYS):
            return False
        try:
            with cache_path.open("rb") as file:
                data = pickle.load(file)
        except (OSError, pickle.UnpicklingError, EOFError):
            logger.warning(f"Ignoring unreadable cache entry: {cache_path.name}")
            return False
        return data.get("schema_version") == self.CACHE_SCHEMA_VERSION
